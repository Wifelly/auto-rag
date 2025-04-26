import hashlib
import json
import os
import sys

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from langchain.docstore.document import Document
from langchain.embeddings import HuggingFaceEmbeddings
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.vectorstores import FAISS
from pptx import Presentation
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn

console = Console()
HASHES_FILE = "file_hashes.json"
INDEX_DIR = "docs_db_index"


def calculate_file_hash(file_path):
    with open(file_path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


def load_hashes():
    if os.path.exists(HASHES_FILE):
        with open(HASHES_FILE) as f:
            return json.load(f)
    return {}


def save_hashes(hashes):
    with open(HASHES_FILE, "w") as f:
        json.dump(hashes, f, indent=2)


def extract_text_from_pdf(file_path):
    text = ""
    with fitz.open(file_path) as doc:
        for page in doc:
            text += page.get_text()
    return text


def extract_text_from_docx(file_path):
    doc = DocxDocument(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text + "\n"
    return text


def extract_texts_with_hash_check(folder_path):
    current_hashes = {}
    full_text = ""
    processed_files = load_hashes()

    try:
        files = os.listdir(folder_path)
    except FileNotFoundError:
        console.log(f"[red]Папка '{folder_path}' не найдена. Завершение работы.[/red]")
        sys.exit(1)

    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        transient=True,
    ) as progress:
        task = progress.add_task("[green]Обработка документов...", total=len(files))

        for filename in files:
            file_path = os.path.join(folder_path, filename)
            if not filename.endswith((".pdf", ".docx", ".pptx")):
                progress.advance(task)
                continue

            file_hash = calculate_file_hash(file_path)
            current_hashes[filename] = file_hash

            if filename in processed_files and processed_files[filename] == file_hash:
                console.log(f"Пропуск {filename} — не изменялся.")
                progress.advance(task)
                continue

            try:
                console.log(f"Обработка {filename}...")
                if filename.endswith(".pdf"):
                    full_text += extract_text_from_pdf(file_path) + "\n"
                elif filename.endswith(".docx"):
                    full_text += extract_text_from_docx(file_path) + "\n"
                elif filename.endswith(".pptx"):
                    full_text += extract_text_from_pptx(file_path) + "\n"
            except Exception as e:
                console.log(f"[red]Ошибка при обработке {filename}: {e}[/red]")

            progress.advance(task)

    save_hashes(current_hashes)
    return full_text


def create_search_db(file_text, knowledge_base_link, chunk_size=1024, chunk_overlap=200, append=True):
    if not file_text.strip():
        console.log("[yellow]Нет новых или изменённых файлов для обработки.[/yellow]")
        return

    console.log("[bold]Разделение текста на чанки...[/bold]")
    splitter = RecursiveCharacterTextSplitter(
        separators=["\n\n", "\n", " "],
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
    chunks = splitter.split_text(file_text)

    source_chunks = []
    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        transient=True,
    ) as progress:
        task = progress.add_task("[cyan]Создание чанков документа...", total=len(chunks))
        for chunkID, chunk in enumerate(chunks):
            source_chunks.append(
                Document(
                    page_content=chunk,
                    metadata={"source": knowledge_base_link, "chunkID": chunkID},
                )
            )
            progress.advance(task)

    if not source_chunks:
        console.log("[red]Не удалось создать базу: нет чанков для обработки.[/red]")
        return

    console.log(f"[green]Создано чанков: {len(source_chunks)}[/green]")

    embedding_model = HuggingFaceEmbeddings(model_name="sentence-transformers/all-MiniLM-L6-v2")

    if append and os.path.exists(os.path.join(INDEX_DIR, "index.faiss")):
        console.log("[yellow]Загрузка существующего индекса FAISS...[/yellow]")
        db = FAISS.load_local(INDEX_DIR, embedding_model)
        db.add_documents(source_chunks)
        db.save_local(INDEX_DIR)
        console.log("[green]Индекс FAISS обновлён новыми документами[/green]")
    else:
        console.log("[blue]Создание новой базы FAISS...[/blue]")
        db = FAISS.from_documents(source_chunks, embedding_model)
        db.save_local(INDEX_DIR)
        console.log("[green]Векторная база успешно создана![/green]")


if __name__ == "__main__":
    console.rule("[bold blue]Обновление базы знаний")
    folder_path = "user_documents"
    new_text = extract_texts_with_hash_check(folder_path)
    create_search_db(new_text, knowledge_base_link="user_documents_kb", append=True)
