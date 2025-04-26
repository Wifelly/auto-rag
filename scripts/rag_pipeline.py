import argparse
import hashlib
import json
import os
import pickle
import sys
import traceback

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from langchain.docstore.document import Document
from langchain.embeddings import HuggingFaceEmbeddings
from langchain.embeddings.base import Embeddings
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.vectorstores import FAISS
from pptx import Presentation
from rich.console import Console
from rich.progress import Progress, SpinnerColumn, TextColumn
from sentence_transformers import SentenceTransformer
from tqdm import tqdm

console = Console()
HASHES_FILE = "file_hashes.json"
USER_DOCS_INDEX_DIR = "docs_db_index"
PUBMED_INDEX_DIR = "pubmed_db_index"


class TqdmEmbeddings(Embeddings):
    def __init__(self, model_name):
        self.model = SentenceTransformer(model_name)

    def embed_documents(self, texts, batch_size=10):
        all_embeddings = []
        for i in tqdm(
            range(0, len(texts), batch_size), desc="Создание эмбеддингов", unit="батч"
        ):
            batch = texts[i : i + batch_size]
            batch_embeddings = self.model.encode(batch, show_progress_bar=False)
            all_embeddings.extend(batch_embeddings)
        return all_embeddings

    def embed_query(self, text):
        return self.model.encode([text])[0]


def calculate_file_hash(file_path):
    with open(file_path, "rb") as f:
        return hashlib.md5(f.read()).hexdigest()


def load_hashes():
    if os.path.exists(HASHES_FILE):
        with open(HASHES_FILE, "r") as f:
            return json.load(f)
    return {}


def save_hashes(hashes):
    with open(HASHES_FILE, "w") as f:
        json.dump(hashes, f, indent=2)


# User documents processing functions
def extract_text_from_pdf(file_path):
    text = ""
    with fitz.open(file_path) as doc:
        for page in doc:
            text += page.get_text()
    return text


def extract_text_from_docx(file_path):
    doc = DocxDocument(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text += shape.text + "\n"
    return text


def extract_texts_with_hash_check(folder_path):
    current_hashes = {}
    full_text = ""
    processed_files = load_hashes()

    try:
        files = os.listdir(folder_path)
    except FileNotFoundError:
        console.log(f"[red]Папка '{folder_path}' не найдена. Завершение работы.[/red]")
        sys.exit(1)

    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        transient=True,
    ) as progress:
        task = progress.add_task("[green]Обработка документов...", total=len(files))

        for filename in files:
            file_path = os.path.join(folder_path, filename)
            if not filename.endswith((".pdf", ".docx", ".pptx")):
                progress.advance(task)
                continue

            file_hash = calculate_file_hash(file_path)
            current_hashes[filename] = file_hash

            if filename in processed_files and processed_files[filename] == file_hash:
                console.log(f"Пропуск {filename} — не изменялся.")
                progress.advance(task)
                continue

            try:
                console.log(f"Обработка {filename}...")
                if filename.endswith(".pdf"):
                    full_text += extract_text_from_pdf(file_path) + "\n"
                elif filename.endswith(".docx"):
                    full_text += extract_text_from_docx(file_path) + "\n"
                elif filename.endswith(".pptx"):
                    full_text += extract_text_from_pptx(file_path) + "\n"
            except Exception as e:
                console.log(f"[red]Ошибка при обработке {filename}: {e}[/red]")

            progress.advance(task)

    save_hashes(current_hashes)
    return full_text


# PubMed text processing functions
def extract_text_from_txt_files(folder_path, max_files=None):
    if not os.path.exists(folder_path):
        console.print(f"[bold red]Папка не найдена:[/bold red] {folder_path}")
        return ""

    txt_files = [
        f
        for f in os.listdir(folder_path)
        if f.endswith(".txt") and os.path.isfile(os.path.join(folder_path, f))
    ]

    if max_files and len(txt_files) > max_files:
        console.print(
            f"[yellow]Ограничение:[/yellow] будет обработано {max_files} из {len(txt_files)} файлов"
        )
        txt_files = txt_files[:max_files]

    if not txt_files:
        console.print("[yellow]В папке нет .txt файлов[/yellow]")
        return ""

    console.print(
        f"[cyan]Обработка {len(txt_files)} .txt файлов из {folder_path}[/cyan]"
    )
    full_text = ""
    current_hashes = {}
    previous_hashes = load_hashes()

    for filename in tqdm(txt_files, desc="Чтение файлов", unit="файл"):
        file_path = os.path.join(folder_path, filename)
        file_hash = calculate_file_hash(file_path)
        current_hashes[filename] = file_hash

        if filename in previous_hashes and previous_hashes[filename] == file_hash:
            console.log(f"Пропуск {filename} — не изменялся.")
            continue

        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                full_text += f.read() + "\n\n"
            console.log(f"Обработан: {filename}")
        except Exception as e:
            console.print(f"[red]Ошибка при чтении {file_path}:[/red] {e}")

    save_hashes(current_hashes)
    return full_text


# Generic vector DB creation function
def create_vector_db(
    text,
    output_dir,
    knowledge_base_link,
    chunk_size=1024,
    chunk_overlap=200,
    model_name="sentence-transformers/all-MiniLM-L6-v2",
    append=False,
    use_tqdm=False,
):
    if not text.strip():
        console.print("[yellow]Нет нового текста для обработки[/yellow]")
        return False

    splitter = RecursiveCharacterTextSplitter(
        separators=["\n\n", "\n", " "],
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )

    console.print("[blue]Разделение текста на чанки...[/blue]")
    chunks = splitter.split_text(text)
    console.print(f"[green]Чанков создано:[/green] {len(chunks)}")

    source_chunks = []

    if use_tqdm:
        # Using tqdm for progress
        for i, chunk in enumerate(
            tqdm(chunks, desc="Подготовка документов", unit="чанк")
        ):
            source_chunks.append(
                Document(
                    page_content=chunk,
                    metadata={"source": knowledge_base_link, "chunk_id": i},
                )
            )
    else:
        # Using rich progress for progress
        with Progress(
            SpinnerColumn(),
            TextColumn("[progress.description]{task.description}"),
            transient=True,
        ) as progress:
            task = progress.add_task(
                "[cyan]Создание чанков документа...", total=len(chunks)
            )
            for i, chunk in enumerate(chunks):
                source_chunks.append(
                    Document(
                        page_content=chunk,
                        metadata={"source": knowledge_base_link, "chunk_id": i},
                    )
                )
                progress.advance(task)

    if not source_chunks:
        console.print("[red]Не удалось создать документы для индексации[/red]")
        return False

    try:
        console.print(f"[bold]Загрузка модели эмбеддингов:[/bold] {model_name}")

        if use_tqdm:
            embeddings = TqdmEmbeddings(model_name)
        else:
            embeddings = HuggingFaceEmbeddings(model_name=model_name)

        os.makedirs(output_dir, exist_ok=True)
        faiss_path = os.path.join(output_dir, "index.faiss")

        if append and os.path.exists(faiss_path):
            console.print("[yellow]Загрузка существующего индекса FAISS...[/yellow]")
            db = FAISS.load_local(output_dir, embeddings)
            db.add_documents(source_chunks)
            db.save_local(output_dir)
            console.print("[green]Индекс FAISS обновлён новыми документами[/green]")
        else:
            console.print("[bold blue]Создание новой базы FAISS...[/bold blue]")
            db = FAISS.from_documents(source_chunks, embeddings)
            db.save_local(output_dir)
            console.print(f"[bold green]База создана в: {output_dir}[/bold green]")

        # If we're processing PubMed data, save metadata separately
        if knowledge_base_link == "pubmed_txt":
            metadocs = []
            for i, chunk in enumerate(chunks):
                metadocs.append(
                    {
                        "title": f"PubMed Chunk {i}",
                        "text": chunk,
                        "source": "pubmed_txt",
                    }
                )

            index_path = os.path.join(output_dir, "index.pkl")
            if append and os.path.exists(index_path):
                with open(index_path, "rb") as f:
                    old_metadocs = pickle.load(f)
                metadocs = old_metadocs + metadocs

            with open(index_path, "wb") as f:
                pickle.dump(metadocs, f)

        return True

    except Exception as e:
        console.print(
            f"[bold red]Ошибка при создании/обновлении индекса:[/bold red] {e}"
        )
        traceback.print_exc()
        return False


def process_user_documents(args):
    console.rule("[bold blue]Обновление базы пользовательских документов")
    folder_path = args.user_docs_dir

    new_text = extract_texts_with_hash_check(folder_path)

    if new_text:
        console.print(
            f"Объём нового текста: {len(new_text)} символов (~{len(new_text.split())} слов)"
        )
        create_vector_db(
            text=new_text,
            output_dir=args.user_docs_output_dir,
            knowledge_base_link="user_documents_kb",
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
            model_name=args.model,
            append=args.append,
            use_tqdm=False,
        )
    else:
        console.print(
            "[bold yellow]Нечего индексировать. Все файлы уже обработаны.[/bold yellow]"
        )


def process_pubmed_documents(args):
    console.rule("[bold blue]Обновление базы PubMed документов")
    folder_path = args.pubmed_dir

    new_text = extract_text_from_txt_files(folder_path, args.max_files)

    if new_text:
        console.print(
            f"Объём нового текста: {len(new_text)} символов (~{len(new_text.split())} слов)"
        )
        create_vector_db(
            text=new_text,
            output_dir=args.pubmed_output_dir,
            knowledge_base_link="pubmed_txt",
            chunk_size=args.chunk_size,
            chunk_overlap=args.chunk_overlap,
            model_name=args.model,
            append=args.append,
            use_tqdm=True,
        )
    else:
        console.print(
            "[bold yellow]Нечего индексировать. Все файлы уже обработаны.[/bold yellow]"
        )


def main():
    parser = argparse.ArgumentParser(
        description="Создание или обновление FAISS-векторной базы из документов"
    )
    parser.add_argument(
        "--user-docs-dir",
        default="user_documents",
        help="Папка с пользовательскими документами",
    )
    parser.add_argument(
        "--user-docs-output-dir",
        default="docs_db_index",
        help="Папка для сохранения индекса пользовательских документов",
    )
    parser.add_argument(
        "--pubmed-dir", default="pubmed_txt", help="Папка с PubMed .txt файлами"
    )
    parser.add_argument(
        "--pubmed-output-dir",
        default="pubmed_db_index",
        help="Папка для сохранения индекса PubMed",
    )
    parser.add_argument(
        "--chunk-size", type=int, default=1024, help="Размер чанка в символах"
    )
    parser.add_argument(
        "--chunk-overlap", type=int, default=200, help="Перекрытие между чанками"
    )
    parser.add_argument(
        "--model",
        default="sentence-transformers/all-MiniLM-L6-v2",
        help="Модель эмбеддингов",
    )
    parser.add_argument(
        "--max-files", type=int, help="Ограничить количество файлов (только для PubMed)"
    )
    parser.add_argument(
        "--append", action="store_true", help="Добавить к существующему индексу"
    )
    parser.add_argument(
        "--only-user-docs",
        action="store_true",
        help="Обработать только пользовательские документы",
    )
    parser.add_argument(
        "--only-pubmed", action="store_true", help="Обработать только PubMed документы"
    )

    args = parser.parse_args()

    # If neither or both flags are specified, process both types
    if (not args.only_user_docs and not args.only_pubmed) or (
        args.only_user_docs and args.only_pubmed
    ):
        process_user_documents(args)
        process_pubmed_documents(args)
    elif args.only_user_docs:
        process_user_documents(args)
    elif args.only_pubmed:
        process_pubmed_documents(args)


if __name__ == "__main__":
    main()
