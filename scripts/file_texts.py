import os

import fitz  # PyMuPDF
from docx import Document as DocxDocument
from pptx import Presentation
from rich.console import Console
from rich.markdown import Markdown

console = Console()


def extract_text_from_pdf(file_path):
    text = ""
    with fitz.open(file_path) as doc:
        for page in doc:
            text += page.get_text()
    return text


def extract_text_from_docx(file_path):
    doc = DocxDocument(file_path)
    return "\n".join([p.text for p in doc.paragraphs])


def extract_text_from_pptx(file_path):
    prs = Presentation(file_path)
    text = ""
    for slide_num, slide in enumerate(prs.slides):
        text += f"\n[Слайд {slide_num + 1}]\n"
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text


def preview_and_save_documents(folder_path, output_folder):
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    files = os.listdir(folder_path)
    supported = (".pdf", ".docx", ".pptx")

    for filename in files:
        if not filename.endswith(supported):
            continue

        file_path = os.path.join(folder_path, filename)
        output_filename = os.path.splitext(filename)[0] + ".txt"
        output_path = os.path.join(output_folder, output_filename)

        if filename.endswith(".pdf"):
            text = extract_text_from_pdf(file_path)
        elif filename.endswith(".docx"):
            text = extract_text_from_docx(file_path)
        elif filename.endswith(".pptx"):
            text = extract_text_from_pptx(file_path)
        else:
            continue

        console.rule(f"[bold green]{filename}")

        if text.strip():
            # Показать только начало текста
            preview = text[:3000] + ("\n... (обрезано)" if len(text) > 3000 else "")
            console.print(Markdown("```\n" + preview + "\n```"))

            # Сохранить в файл
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(text)
            console.log(f"[cyan]Текст сохранён в:[/cyan] {output_path}")
        else:
            console.print("[italic yellow]Файл не содержит текста или не удалось извлечь содержимое.[/italic]")


if __name__ == "__main__":
    input_folder = "user_documents"
    output_folder = "texts"
    preview_and_save_documents(input_folder, output_folder)
