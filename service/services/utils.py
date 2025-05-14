# --- utils.py ---

import asyncio
import json
import re
from concurrent.futures import ThreadPoolExecutor
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import aiofiles
import httpx

from service.monitoring.logger import logger

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx")
DEFAULT_HASHES_FILE = "file_hashes.json"
MAX_TEXT_SIZE = 50_000_000  # 50MB


class Utils:
    executor = ThreadPoolExecutor()

    def __init__(self, hashes_file: str = DEFAULT_HASHES_FILE) -> None:
        self.hashes_file: Path = Path(hashes_file)

    @staticmethod
    def clean_text(text: str) -> str:
        text = text.replace("\x00", "")
        text = re.sub(r"\s+", " ", text)
        return text.strip()

    @staticmethod
    def calculate_file_hash(file_path: Path) -> str:
        with file_path.open("rb") as f:
            return sha256(f.read()).hexdigest()

    def load_hashes(self) -> dict[str, str]:
        if self.hashes_file.exists():
            try:
                with self.hashes_file.open("r", encoding="utf-8") as f:
                    return json.load(f)
            except json.JSONDecodeError:
                logger.error("Ошибка чтения файла хешей. Начинаем заново.")
        return {}

    def save_hashes(self, hashes: dict[str, str]) -> None:
        with self.hashes_file.open("w", encoding="utf-8") as f:
            json.dump(hashes, f, indent=2, sort_keys=True)

    @staticmethod
    async def extract_text(file_path: Path) -> str:
        ext = file_path.suffix.lower()
        loop = asyncio.get_running_loop()

        if ext == ".pdf":
            return await loop.run_in_executor(Utils.executor, Utils._extract_pdf_text, file_path)
        if ext == ".docx":
            return await loop.run_in_executor(Utils.executor, Utils._extract_docx_text, file_path)
        if ext == ".pptx":
            return await loop.run_in_executor(Utils.executor, Utils._extract_pptx_text, file_path)

        raise ValueError(f"Неподдерживаемое расширение файла: {ext}")

    @staticmethod
    def _extract_pdf_text(file_path: Path) -> str:
        from fitz import open as fitz_open

        doc = fitz_open(file_path)
        return "\n".join(page.get_text() for page in doc)

    @staticmethod
    def _extract_docx_text(file_path: Path) -> str:
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    @staticmethod
    def _extract_pptx_text(file_path: Path) -> str:
        from pptx import Presentation

        prs = Presentation(file_path)
        return "\n".join(
            shape.text
            for slide in prs.slides
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )

    @staticmethod
    async def async_read_txt_file(file_path: Path) -> str:
        try:
            async with aiofiles.open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return await f.read()
        except Exception as e:
            raise OSError(f"Ошибка чтения TXT файла: {file_path}") from e

    @staticmethod
    async def download_xml(url: str) -> str:
        async with httpx.AsyncClient() as client:
            response = await client.get(url)
            response.raise_for_status()
            return response.text

    @staticmethod
    async def extract_text_from_bytes(filename: str, content: bytes) -> str:
        from docx import Document as DocxDocument
        from fitz import open as fitz_open
        from pptx import Presentation

        ext = Path(filename).suffix.lower()

        try:
            if ext == ".docx":
                doc = DocxDocument(BytesIO(content))
                return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

            elif ext == ".pptx":
                prs = Presentation(BytesIO(content))
                return "\n".join(
                    shape.text
                    for slide in prs.slides
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                )

            elif ext == ".pdf":
                doc = fitz_open(stream=content, filetype="pdf")
                return "\n".join(page.get_text() for page in doc)

            elif ext == ".txt":
                return content.decode("utf-8", errors="ignore")

            else:
                raise ValueError(f"Неподдерживаемый формат файла: {ext}")
        except Exception as e:
            logger.error(f"Ошибка обработки файла {filename}: {e}")
            return ""


__all__ = ["Utils", "SUPPORTED_EXTENSIONS", "MAX_TEXT_SIZE"]
